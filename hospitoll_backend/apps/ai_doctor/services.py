import io
import os
import logging
from django.conf import settings

logger = logging.getLogger(__name__)

_PROMPT_UZ = (
    "Siz tibbiy AI yordamchisiz. Berilgan laboratoriya yoki tibbiy tahlil "
    "natijalarini quyidagi strukturada oddiy va tushunarli o'zbek tilida izohlab bering:\n\n"
    
    "ALOHIDA USULDA TAHLIL QILING:\n"
    "1. DASTLAB: Normal ko'rsatkichlar vs ME'YORDAN CHIQQANLAR ni solishtiring\n"
    "2. AGAR HAMMASI NORMAL: 'Sog'lig'ingiz a'lo darajada' degan xabarni BIRINCHI qatorga qo'ying\n"
    "3. AGAR MUAMMO BOR: ME'YORDAN CHIQQAN ko'rsatkichlarni AVVAL keltirib, keyin qolganlarini yozing\n\n"
    
    "NATIJA FORMATI:\n"
    "--- QISQA XULOSA ---\n"
    "[Agar yaxshi bo'lsa]: Sog'lig'ingiz a'lo darajada. [Qolgan ma'lumotlar]\n"
    "[Agar muammoli bo'lsa]: [Me'yordan chiqqan ko'rsatkichlarni qisqacha]\n\n"
    
    "--- MUHIM TOPILMALAR ---\n"
    "[Faqat me'yordan chiqqan/xavfli ko'rsatkichlarni sanab o'ting]\n\n"
    
    "--- NIMA QILISH KERAK ---\n"
    "[Tavsiyalar keyin]\n\n"
    
    "ESLATMA: Bu yakuniy diagnoz emasligini va shifokorga murojaat qilish kerakligini ayting."
)

_IMAGE_MIME = {
    "image": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "png": "image/png",
    "gif": "image/gif",
    "webp": "image/webp",
    "bmp": "image/bmp",
}

_ANALYSIS_KEYWORDS = (
    "analiz",
    "tahlil",
    "hemoglobin",
    "leykotsit",
    "eritrotsit",
    "glyukoza",
    "cholesterol",
    "triglyceride",
    "hba1c",
    "creatinine",
    "bilirubin",
    "ast",
    "alt",
    "mcv",
    "mch",
    "plt",
    "wbc",
    "rbc",
)

_NOT_ANALYSIS_MESSAGE = "Bu yuklangan narsa boshqa, bunda analiz ma'lumotlari yo'q."


class NonAnalysisContentError(ValueError):
    """Raised when uploaded content does not look like a medical/lab analysis."""


def _raise_if_non_analysis_text(text: str) -> None:
    content = (text or "").strip().lower()
    if len(content) < 40:
        raise NonAnalysisContentError(_NOT_ANALYSIS_MESSAGE)

    has_keyword = any(k in content for k in _ANALYSIS_KEYWORDS)
    has_digit = any(ch.isdigit() for ch in content)
    if not has_keyword or not has_digit:
        raise NonAnalysisContentError(_NOT_ANALYSIS_MESSAGE)


def analyze_lab_image_uzbek(image_bytes: bytes, file_type: str = "image") -> str:
    api_key = getattr(settings, "GEMINI_API_KEY", "") or os.getenv("GEMINI_API_KEY", "")
    model_name = getattr(settings, "GEMINI_VISION_MODEL", "") or os.getenv("GEMINI_VISION_MODEL", "gemini-2.5-flash")

    if not api_key:
        raise ValueError("GEMINI_API_KEY sozlanmagan. Tizim administratoriga murojaat qiling.")

    try:
        from google import genai
        from google.genai import types
    except ImportError:
        raise ValueError("google-genai kutubxonasi o'rnatilmagan.")

    client = genai.Client(api_key=api_key)

    if file_type == "pdf":
        contents = [
            types.Part.from_bytes(data=image_bytes, mime_type="application/pdf"),
            _PROMPT_UZ
            + "\n\nMUHIM: Agar yuklangan fayl tibbiy/laboratoriya analiz hujjati bo'lmasa,"
            + " faqat quyidagi bitta jumlani qaytaring: "
            + _NOT_ANALYSIS_MESSAGE,
        ]
    elif file_type in ("excel", "word", "html"):
        extracted = _extract_text(image_bytes, file_type)
        _raise_if_non_analysis_text(extracted)
        contents = [_PROMPT_UZ + "\n\nTahlil ma'lumotlari:\n" + extracted]
    else:
        mime = _IMAGE_MIME.get(file_type, "image/png")
        contents = [
            types.Part.from_bytes(data=image_bytes, mime_type=mime),
            _PROMPT_UZ
            + "\n\nMUHIM: Agar rasm tibbiy/laboratoriya analiz rasmi bo'lmasa,"
            + " faqat quyidagi bitta jumlani qaytaring: "
            + _NOT_ANALYSIS_MESSAGE,
        ]

    response = client.models.generate_content(model=model_name, contents=contents)
    result = (response.text or "").strip()
    if not result:
        raise ValueError("AI dan bo'sh javob qaytdi.")
    if _NOT_ANALYSIS_MESSAGE in result:
        raise NonAnalysisContentError(_NOT_ANALYSIS_MESSAGE)
    return result


def _extract_text(file_bytes: bytes, file_type: str) -> str:
    if file_type == "excel":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"--- Varaq: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) if c is not None else "" for c in row]
                    if any(v.strip() for v in cells):
                        lines.append("\t".join(cells))
            return "\n".join(lines)
        except Exception as e:
            logger.error("Error reading excel file: %s", e)
            return ""

    if file_type == "word":
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            logger.error("Error reading docx file: %s", e)
            return ""

    if file_type == "html":
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(file_bytes, "html.parser")
            return soup.get_text(separator="\n", strip=True)
        except Exception as e:
            logger.error("Error reading html file: %s", e)
            return ""

    return ""
